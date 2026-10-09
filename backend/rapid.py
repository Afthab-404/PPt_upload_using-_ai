
from pptx import Presentation
from pathlib import Path

template_dir = Path("templates")

files = list(template_dir.glob("*.pptx"))

if not files:
    print("No .pptx files found in the templates folder.")
    print("Upload your template first.")
    raise SystemExit(1)

print("\nAvailable templates:")
for index, file in enumerate(files, start=1):
    print(f"{index}. {file.name}")

template_path = files[-1]
prs = Presentation(str(template_path))

print(f"\nInspecting: {template_path}")
print(f"Existing slides: {len(prs.slides)}")
print(f"Available slide layouts: {len(prs.slide_layouts)}")

for index, layout in enumerate(prs.slide_layouts):
    print(f"\nLayout {index}: {layout.name}")

    for placeholder in layout.placeholders:
        print(
            f"  Placeholder {placeholder.placeholder_format.idx}: "
            f"{placeholder.name}"
        )

for index, slide in enumerate(prs.slides, start=1):
    print(f"\nSlide {index}")
    print(f"Layout: {slide.slide_layout.name}")

    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            print(f"  Text: {shape.text[:150]!r}")

        if shape.is_placeholder:
            print(f"  Placeholder: {shape.name}")
